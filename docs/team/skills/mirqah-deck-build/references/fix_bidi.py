"""Post-process: mark Arabic runs as ar-SA (+cs font) and split Latin tokens into en-US runs."""
import sys, re, copy
from pptx import Presentation
from pptx.oxml.ns import qn
LAT = re.compile(r'[A-Za-z][A-Za-z0-9.\-_/:]*(?:[ ][A-Za-z0-9][A-Za-z0-9.\-_/:]*)*|[0-9]+(?:\.[0-9]+)?(?=[ ]?[A-Za-z])')
def setlang(r, lang):
    rPr = r.find(qn('a:rPr'))
    if rPr is None:
        rPr = r.makeelement(qn('a:rPr'), {}); r.insert(0, rPr)
    rPr.set('lang', lang)
    if lang == 'ar-SA':
        rPr.set('altLang', 'en-US')
        lat = rPr.find(qn('a:latin'))
        if lat is not None and rPr.find(qn('a:cs')) is None:
            cs = rPr.makeelement(qn('a:cs'), {'typeface': lat.get('typeface')}); lat.addnext(cs)
def fix(path_in, path_out):
    p = Presentation(path_in); n = 0
    for s in p.slides:
        for sh in s.shapes:
            if not sh.has_text_frame: continue
            for para in sh.text_frame.paragraphs:
                for r in list(para._p.findall(qn('a:r'))):
                    t = r.find(qn('a:t')).text or ''
                    if not re.search('[\u0600-\u06FF]', t):
                        continue
                    # split into Arabic / Latin segments
                    parts, last = [], 0
                    for m in re.finditer(r'[A-Za-z0-9][A-Za-z0-9.\-_/:]*(?: [A-Za-z0-9][A-Za-z0-9.\-_/:]*)*', t):
                        if not re.search('[A-Za-z]', m.group()) and not re.search('[A-Za-z]', t[max(0,m.start()-1):m.end()+1]):
                            continue
                        parts.append((t[last:m.start()], 'ar-SA')); parts.append((m.group(), 'en-US')); last = m.end()
                    parts.append((t[last:], 'ar-SA'))
                    parts = [x for x in parts if x[0]]
                    prev = r
                    for i, (txt, lang) in enumerate(parts):
                        nr = r if i == 0 else copy.deepcopy(r)
                        nr.find(qn('a:t')).text = txt
                        setlang(nr, lang)
                        if i: prev.addnext(nr)
                        prev = nr; n += 1
    p.save(path_out); print("runs set:", n)
if __name__ == '__main__':
    fix(sys.argv[1], sys.argv[2])
